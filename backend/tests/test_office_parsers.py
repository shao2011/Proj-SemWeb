import re
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from rag_ingestion.models import BlockType, CellRangeLocator, SourceType
from rag_ingestion.ingestion.parsers.office import parse_docx, parse_pptx, parse_xlsx
from rag_ingestion.storage.media import LocalMediaStore

JPEG = b"\xff\xd8\xff\xe0" + b"JFIF" + b"\x00" * 64
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64

_DRAWING = """<?xml version="1.0" encoding="UTF-8"?>
<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
          xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <xdr:twoCellAnchor>
    <xdr:from><xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>0</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>
    <xdr:to><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>5</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
    <xdr:pic>
      <xdr:nvPicPr><xdr:cNvPr id="2" name="QR code"/><xdr:cNvPicPr/></xdr:nvPicPr>
      <xdr:blipFill>
        <a:blip xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:embed="rId1"/>
      </xdr:blipFill>
    </xdr:pic>
  </xdr:twoCellAnchor>
</xdr:wsDr>
"""

_DRAWING_RELS = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1"
                Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
                Target="../media/{media_name}"/>
</Relationships>
"""

_SHEET_RELS = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1"
                Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"
                Target="../drawings/drawing1.xml"/>
</Relationships>
"""


def _write_zip(path: Path, files: dict[str, str]) -> None:
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)


def _repack(path: Path, additions: dict[str, bytes | str], patch: dict[str, str]) -> None:
    """Rewrite an xlsx in place with extra parts and patched sheet xml.

    A zip entry cannot be replaced in place, so the whole archive is rebuilt.
    """
    with ZipFile(path) as archive:
        existing = {name: archive.read(name) for name in archive.namelist()}
    for name, replacement in patch.items():
        existing[name] = replacement.encode()
    existing.update(
        {
            name: content if isinstance(content, bytes) else content.encode()
            for name, content in additions.items()
        }
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in existing.items():
            archive.writestr(name, content)


def _cache_formula_results(path: Path, results: dict[str, str]) -> None:
    """Give formula cells the cached result Excel would have stored.

    openpyxl saves formulas with no cached value, so a workbook it wrote looks
    exactly like one Excel has never opened. These tests need the other case.
    """
    with ZipFile(path) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode()
    for coordinate, value in results.items():
        sheet_xml = re.sub(
            rf'(<c r="{coordinate}"[^>]*>)(<f>[^<]*</f>)',
            rf"\1\2<v>{value}</v>",
            sheet_xml,
        )
    _repack(path, {}, {"xl/worksheets/sheet1.xml": sheet_xml})


def _embed_picture(path: Path, payload: bytes, media_name: str) -> None:
    _repack(
        path,
        {
            f"xl/media/{media_name}": payload,
            "xl/drawings/drawing1.xml": _DRAWING,
            "xl/drawings/_rels/drawing1.xml.rels": _DRAWING_RELS.format(media_name=media_name),
            "xl/worksheets/_rels/sheet1.xml.rels": _SHEET_RELS,
        },
        {},
    )


def _row_text(result, row_number: int) -> str:
    return next(
        block.text
        for block in result.blocks
        if block.type is BlockType.TABLE_ROW and block.metadata.get("row") == row_number
    )


def test_docx_preserves_heading_hierarchy_and_table_cells(tmp_path: Path) -> None:
    path = tmp_path / "policy.docx"
    _write_zip(
        path,
        {
            "word/document.xml": """
                <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                  <w:body>
                    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Policy</w:t></w:r></w:p>
                    <w:p><w:r><w:t>Applies to every team.</w:t></w:r></w:p>
                    <w:tbl><w:tr>
                      <w:tc><w:p><w:r><w:t>Owner</w:t></w:r></w:p></w:tc>
                      <w:tc><w:p><w:r><w:t>Legal</w:t></w:r></w:p></w:tc>
                    </w:tr></w:tbl>
                  </w:body>
                </w:document>
            """,
        },
    )

    result = parse_docx(
        path,
        owner_user_id="user-a",
        project_id="project-a",
        source_uri="upload://policy.docx",
    )

    assert result.document.source_type is SourceType.DOCX
    assert [block.type for block in result.blocks] == [
        BlockType.HEADING,
        BlockType.PARAGRAPH,
        BlockType.TABLE,
        BlockType.TABLE_ROW,
        BlockType.TABLE_CELL,
        BlockType.TABLE_CELL,
    ]
    heading, paragraph = result.blocks[:2]
    assert heading.section_path == ["Policy"]
    assert paragraph.section_path == ["Policy"]
    assert paragraph.parent_block_id == heading.block_id
    assert all(block.confidence == 1.0 for block in result.blocks)


def _ledger(tmp_path: Path, name: str = "ledger.xlsx") -> Path:
    path = tmp_path / name
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Revenue"
    sheet["A2"] = "Amount"
    sheet["B2"] = "=SUM(B3:B4)"
    sheet["B3"] = 10
    sheet["B4"] = 20
    workbook.save(path)
    return path


def _parse(path: Path, **extra):
    return parse_xlsx(
        path,
        owner_user_id="user-a",
        project_id="project-a",
        source_uri=f"upload://{path.name}",
        **extra,
    )


def test_xlsx_preserves_formula_merged_range_and_cell_locators(tmp_path: Path) -> None:
    path = _ledger(tmp_path)

    result = _parse(path)

    assert result.document.source_type is SourceType.XLSX
    revenue = next(
        block
        for block in result.blocks
        if block.type is BlockType.TABLE_CELL and block.text == "Revenue"
    )
    formula = next(
        block
        for block in result.blocks
        if block.type is BlockType.TABLE_CELL and block.metadata.get("formula")
    )
    assert revenue.metadata["merged_range"] == "A1:B1"
    assert revenue.locators == [CellRangeLocator(sheet="Data", range="A1:B1")]
    assert formula.metadata["formula"] == "=SUM(B3:B4)"
    assert formula.locators == [CellRangeLocator(sheet="Data", range="B2")]
    assert all(block.confidence == 1.0 for block in result.blocks)


def test_xlsx_rows_show_the_computed_value_not_the_formula(tmp_path: Path) -> None:
    """The bug this replaces: a reader saw `=SUM(B3:B4)` where 30 belonged."""
    path = _ledger(tmp_path)
    _cache_formula_results(path, {"B2": "30"})

    result = _parse(path)

    assert "30" in _row_text(result, 2)
    assert "SUM" not in _row_text(result, 2)
    formula = next(
        block
        for block in result.blocks
        if block.type is BlockType.TABLE_CELL and block.metadata.get("formula")
    )
    assert formula.text == "30"
    assert formula.metadata["formula"] == "=SUM(B3:B4)"


def test_xlsx_falls_back_to_the_formula_when_nothing_was_ever_computed(
    tmp_path: Path,
) -> None:
    """A workbook Excel never opened caches no results at all.

    Dropping those cells would empty the sheet, which is strictly worse than
    showing the formula, so the formula stays as the last resort.
    """
    path = _ledger(tmp_path)

    result = _parse(path)

    assert "=SUM(B3:B4)" in _row_text(result, 2)


def test_xlsx_leaves_a_calculated_blank_blank_instead_of_showing_its_formula(
    tmp_path: Path,
) -> None:
    """Excel caches nothing for a formula whose result was empty.

    Deciding the fallback per cell put `=IF(...)` back into every cell that had
    correctly evaluated to blank, which is most of a real sheet.
    """
    path = tmp_path / "blanks.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = 10
    sheet["B1"] = "=A1*2"
    sheet["C1"] = '=IF(A1=0,A1,"")'
    workbook.save(path)
    _cache_formula_results(path, {"B1": "20"})

    result = _parse(path)

    assert "20" in _row_text(result, 1)
    assert "IF(" not in _row_text(result, 1)


def test_xlsx_keeps_empty_cells_so_a_value_stays_under_its_column(
    tmp_path: Path,
) -> None:
    """Collapsing blanks shifted every later value one column left."""
    path = tmp_path / "gaps.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "Hùng"
    sheet["B1"] = "Ngọc"
    sheet["C1"] = "Đạt"
    sheet["A2"] = 10
    sheet["C2"] = 30
    workbook.save(path)

    result = _parse(path)

    assert _row_text(result, 1).split("|")[2:5] == [" Hùng ", " Ngọc ", " Đạt "]
    assert _row_text(result, 2).split("|")[2:5] == [" 10 ", "  ", " 30 "]


def test_xlsx_gives_every_sheet_its_own_heading(tmp_path: Path) -> None:
    """Two sheets both numbered their rows from one and collided in the output."""
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Công nợ"
    first["A1"] = "owed"
    second = workbook.create_sheet("QR")
    second["A1"] = "code"
    workbook.save(path)

    result = _parse(path)

    headings = [
        block.text for block in result.blocks if block.type is BlockType.HEADING
    ]
    assert headings == ["Công nợ", "QR"]
    rows = [block for block in result.blocks if block.type is BlockType.TABLE_ROW]
    assert {block.section_path[0] for block in rows} == {"Công nợ", "QR"}


def test_xlsx_drops_rows_whose_formulas_all_resolved_to_nothing(tmp_path: Path) -> None:
    """Fill-down tails are the majority of a real sheet and carry no content."""
    path = tmp_path / "tail.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "Bún chả"
    sheet["B1"] = 38
    sheet["B2"] = "=SUM(A2:A2)"
    workbook.save(path)
    _cache_formula_results(path, {"B2": "0"})

    result = _parse(path)

    rows = [block.metadata["row"] for block in result.blocks if block.type is BlockType.TABLE_ROW]
    assert rows == [1]


def test_xlsx_renders_dates_and_whole_numbers_the_way_a_reader_writes_them(
    tmp_path: Path,
) -> None:
    path = tmp_path / "dates.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = datetime(2026, 6, 26)
    sheet["B1"] = 38.0
    sheet["C1"] = 3.75
    workbook.save(path)

    result = _parse(path)

    assert "2026-06-26" in _row_text(result, 1)
    assert "00:00:00" not in _row_text(result, 1)
    assert "| 38 |" in _row_text(result, 1)
    assert "3.75" in _row_text(result, 1)


def test_xlsx_stores_an_embedded_picture_and_points_a_figure_at_it(
    tmp_path: Path,
) -> None:
    path = _ledger(tmp_path)
    _embed_picture(path, PNG, "image1.png")
    media = LocalMediaStore(tmp_path / "media")

    result = _parse(path, media_store=media)

    figure = next(block for block in result.blocks if block.type is BlockType.FIGURE)
    assert figure.media_ref is not None
    assert media.get(figure.media_ref) == PNG
    assert figure.locators == [CellRangeLocator(sheet="Data", range="A1:D6")]
    assert "QR code" in figure.text


def test_xlsx_names_a_picture_by_its_content_not_its_stored_extension(
    tmp_path: Path,
) -> None:
    """This workbook ships a JPEG called image4.tmp; the extension is a lie."""
    path = _ledger(tmp_path)
    _embed_picture(path, JPEG, "image1.tmp")
    media = LocalMediaStore(tmp_path / "media")

    result = _parse(path, media_store=media)

    figure = next(block for block in result.blocks if block.type is BlockType.FIGURE)
    assert figure.media_ref.endswith(".jpg")


def test_xlsx_without_a_media_store_parses_the_cells_and_skips_the_pictures(
    tmp_path: Path,
) -> None:
    path = _ledger(tmp_path)
    _embed_picture(path, PNG, "image1.png")

    result = _parse(path)

    assert not [block for block in result.blocks if block.type is BlockType.FIGURE]
    assert any(block.type is BlockType.TABLE_ROW for block in result.blocks)


_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def _shape(*lines: str, title: bool = False) -> str:
    placeholder = f'<p:nvSpPr><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>' if title else ""
    paragraphs = "".join(
        f"<a:p><a:r><a:t>{line}</a:t></a:r></a:p>" for line in lines
    )
    return f"<p:sp>{placeholder}<p:txBody>{paragraphs}</p:txBody></p:sp>"


def _picture(name: str, rel_id: str = "rId9") -> str:
    return (
        f'<p:pic><p:nvPicPr><p:cNvPr id="9" name="{name}"/></p:nvPicPr>'
        f'<p:blipFill><a:blip r:embed="{rel_id}"/></p:blipFill></p:pic>'
    )


def _table(rows: list[list[str]]) -> str:
    body = "".join(
        "<a:tr>"
        + "".join(f"<a:tc><a:txBody><a:p><a:r><a:t>{cell}</a:t></a:r></a:p></a:txBody></a:tc>" for cell in row)
        + "</a:tr>"
        for row in rows
    )
    return f"<p:graphicFrame><a:graphic><a:graphicData><a:tbl>{body}</a:tbl></a:graphicData></a:graphic></p:graphicFrame>"


def _slide(*shapes: str, hidden: bool = False) -> str:
    show = ' show="0"' if hidden else ""
    return (
        f'<p:sld xmlns:p="{_P}" xmlns:a="{_A}" xmlns:r="{_R}"{show}>'
        f"<p:cSld><p:spTree>{''.join(shapes)}</p:spTree></p:cSld></p:sld>"
    )


def _notes(*lines: str, slide_number: str | None = None) -> str:
    paragraphs = "".join(f"<a:p><a:r><a:t>{line}</a:t></a:r></a:p>" for line in lines)
    if slide_number is not None:
        paragraphs += f'<a:p><a:fld id="x" type="slidenum"><a:t>{slide_number}</a:t></a:fld></a:p>'
    return (
        f'<p:notes xmlns:p="{_P}" xmlns:a="{_A}"><p:cSld><p:spTree>'
        f"<p:sp><p:txBody>{paragraphs}</p:txBody></p:sp>"
        f"</p:spTree></p:cSld></p:notes>"
    )


def _rels(*entries: tuple[str, str, str]) -> str:
    body = "".join(
        f'<Relationship Id="{rel_id}" Type="{_R}/{kind}" Target="{target}"/>'
        for rel_id, kind, target in entries
    )
    return f'<Relationships xmlns="{_PKG}">{body}</Relationships>'


def _presentation(*rel_ids: str) -> dict[str, str]:
    """A deck whose slide order is given by `rel_ids`, not by part names."""
    slide_ids = "".join(
        f'<p:sldId id="{256 + position}" r:id="{rel_id}"/>'
        for position, rel_id in enumerate(rel_ids)
    )
    return {
        "ppt/presentation.xml": (
            f'<p:presentation xmlns:p="{_P}" xmlns:r="{_R}">'
            f"<p:sldIdLst>{slide_ids}</p:sldIdLst></p:presentation>"
        ),
        "ppt/_rels/presentation.xml.rels": _rels(
            *(
                (rel_id, "slide", f"slides/slide{index}.xml")
                for index, rel_id in enumerate(sorted(rel_ids), start=1)
            )
        ),
    }


def _deck(tmp_path: Path, files: dict[str, str], media: dict[str, bytes] | None = None) -> Path:
    path = tmp_path / "briefing.pptx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
        for name, payload in (media or {}).items():
            archive.writestr(name, payload)
    return path


def _parse_deck(path: Path, **extra):
    return parse_pptx(
        path,
        owner_user_id="user-a",
        project_id="project-a",
        source_uri="upload://briefing.pptx",
        **extra,
    )


def test_pptx_reads_slides_in_presentation_order_not_file_name_order(
    tmp_path: Path,
) -> None:
    """Reordering a deck rewrites `sldIdLst` and never renames a part."""
    files = {
        **_presentation("rId3", "rId1", "rId2"),
        "ppt/slides/slide1.xml": _slide(_shape("Doanh thu", title=True)),
        "ppt/slides/slide2.xml": _slide(_shape("Kết luận", title=True)),
        "ppt/slides/slide3.xml": _slide(_shape("Mở đầu", title=True)),
    }

    result = _parse_deck(_deck(tmp_path, files))

    titles = [block.text for block in result.blocks if block.type is BlockType.SLIDE]
    assert titles == ["Mở đầu", "Doanh thu", "Kết luận"]
    assert [block.metadata["slide"] for block in result.blocks if block.type is BlockType.SLIDE] == [1, 2, 3]


def test_pptx_attaches_notes_through_the_slide_relationship(tmp_path: Path) -> None:
    """The notes part number is creation order, not slide order.

    Guessing `notesSlide{n}.xml` handed slide 2's notes to slide 1, which is a
    silent wrong answer rather than a visible failure.
    """
    files = {
        **_presentation("rId1", "rId2"),
        "ppt/slides/slide1.xml": _slide(_shape("Doanh thu", title=True)),
        "ppt/slides/slide2.xml": _slide(_shape("Kiến trúc", title=True)),
        "ppt/slides/_rels/slide2.xml.rels": _rels(
            ("rId2", "notesSlide", "../notesSlides/notesSlide1.xml")
        ),
        "ppt/notesSlides/notesSlide1.xml": _notes("Nói kỹ phần kiến trúc"),
    }

    result = _parse_deck(_deck(tmp_path, files))

    notes = [
        block for block in result.blocks if block.metadata.get("speaker_notes") is True
    ]
    assert len(notes) == 1
    assert notes[0].metadata["slide"] == 2
    assert "Nói kỹ phần kiến trúc" in notes[0].text


def test_pptx_keeps_speaker_notes_out_of_the_slide_body(tmp_path: Path) -> None:
    """A reader could not tell the spoken script from what was on the screen."""
    files = {
        **_presentation("rId1"),
        "ppt/slides/slide1.xml": _slide(_shape("Doanh thu", title=True), _shape("Tăng 12%")),
        "ppt/slides/_rels/slide1.xml.rels": _rels(
            ("rId2", "notesSlide", "../notesSlides/notesSlide1.xml")
        ),
        "ppt/notesSlides/notesSlide1.xml": _notes("Đừng đọc số liệu"),
    }

    result = _parse_deck(_deck(tmp_path, files))

    body = [
        block.text
        for block in result.blocks
        if block.metadata.get("speaker_notes") is not True
    ]
    assert not any("Đừng đọc số liệu" in text for text in body)
    notes = next(b for b in result.blocks if b.metadata.get("speaker_notes") is True)
    assert notes.type is BlockType.SPEECH_TURN


def test_pptx_drops_the_slide_number_placeholder_from_the_notes(tmp_path: Path) -> None:
    files = {
        **_presentation("rId1"),
        "ppt/slides/slide1.xml": _slide(_shape("Doanh thu", title=True)),
        "ppt/slides/_rels/slide1.xml.rels": _rels(
            ("rId2", "notesSlide", "../notesSlides/notesSlide1.xml")
        ),
        "ppt/notesSlides/notesSlide1.xml": _notes("Mở đầu nhẹ nhàng", slide_number="7"),
    }

    result = _parse_deck(_deck(tmp_path, files))

    notes = next(b for b in result.blocks if b.metadata.get("speaker_notes") is True)
    assert "7" not in notes.text


def test_pptx_gives_every_slide_its_own_section(tmp_path: Path) -> None:
    """With one section for a whole deck, routing cannot narrow to a slide."""
    files = {
        **_presentation("rId1", "rId2"),
        "ppt/slides/slide1.xml": _slide(_shape("Doanh thu", title=True), _shape("Tăng 12%")),
        "ppt/slides/slide2.xml": _slide(_shape("Kiến trúc", title=True), _shape("Ba tầng")),
    }

    result = _parse_deck(_deck(tmp_path, files))

    assert {tuple(block.section_path) for block in result.blocks} == {
        ("Doanh thu",),
        ("Kiến trúc",),
    }


def test_pptx_splits_a_slide_into_its_own_blocks(tmp_path: Path) -> None:
    """One block per slide made the whole slide a markdown heading, and made the
    body too big for the chunker to ever split."""
    files = {
        **_presentation("rId1"),
        "ppt/slides/slide1.xml": _slide(
            _shape("Kiến trúc", title=True), _shape("Tầng tiếp nhận", "Tầng xử lý")
        ),
    }

    result = _parse_deck(_deck(tmp_path, files))

    assert [(block.type, block.text) for block in result.blocks] == [
        (BlockType.SLIDE, "Kiến trúc"),
        (BlockType.PARAGRAPH, "Tầng tiếp nhận"),
        (BlockType.PARAGRAPH, "Tầng xử lý"),
    ]


def test_pptx_reads_shapes_inside_a_group(tmp_path: Path) -> None:
    """A diagram is a group of shapes, which is most of what a deck's meaning
    lives in. Reading only the top level of the shape tree loses all of it."""
    group = f"<p:grpSp>{_shape('Receipt')}{_shape('S3')}{_shape('extraction')}</p:grpSp>"
    files = {
        **_presentation("rId1"),
        "ppt/slides/slide1.xml": _slide(_shape("Overview", title=True), group),
    }

    result = _parse_deck(_deck(tmp_path, files))

    assert [
        block.text for block in result.blocks if block.type is BlockType.PARAGRAPH
    ] == ["Receipt", "S3", "extraction"]


def test_pptx_reads_a_picture_nested_in_a_group(tmp_path: Path) -> None:
    group = f'<p:grpSp>{_picture("Sơ đồ luồng")}</p:grpSp>'
    files = {
        **_presentation("rId1"),
        "ppt/slides/slide1.xml": _slide(_shape("Kiến trúc", title=True), group),
        "ppt/slides/_rels/slide1.xml.rels": _rels(("rId9", "image", "../media/image1.png")),
    }
    media = LocalMediaStore(tmp_path / "media")

    result = _parse_deck(
        _deck(tmp_path, files, {"ppt/media/image1.png": PNG}), media_store=media
    )

    figure = next(block for block in result.blocks if block.type is BlockType.FIGURE)
    assert media.get(figure.media_ref) == PNG


def test_pptx_renders_a_slide_table_as_a_table(tmp_path: Path) -> None:
    files = {
        **_presentation("rId1"),
        "ppt/slides/slide1.xml": _slide(
            _shape("Chi phí", title=True),
            _table([["Hạng mục", "Tiền"], ["Máy chủ", "120"]]),
        ),
    }

    result = _parse_deck(_deck(tmp_path, files))

    header = next(block for block in result.blocks if block.type is BlockType.TABLE)
    rows = [block for block in result.blocks if block.type is BlockType.TABLE_ROW]
    assert header.text == "| Hạng mục | Tiền |\n|---|---|"
    assert [block.text for block in rows] == ["| Máy chủ | 120 |"]
    assert all(block.parent_block_id == header.block_id for block in rows)


def test_pptx_stores_a_slide_picture_and_points_a_figure_at_it(tmp_path: Path) -> None:
    files = {
        **_presentation("rId1"),
        "ppt/slides/slide1.xml": _slide(_shape("Kiến trúc", title=True), _picture("Sơ đồ luồng")),
        "ppt/slides/_rels/slide1.xml.rels": _rels(("rId9", "image", "../media/image1.png")),
    }
    media = LocalMediaStore(tmp_path / "media")

    result = _parse_deck(
        _deck(tmp_path, files, {"ppt/media/image1.png": PNG}), media_store=media
    )

    figure = next(block for block in result.blocks if block.type is BlockType.FIGURE)
    assert media.get(figure.media_ref) == PNG
    assert "Sơ đồ luồng" in figure.text
    assert figure.section_path == ["Kiến trúc"]


def test_pptx_names_a_slide_picture_by_its_content_not_its_extension(
    tmp_path: Path,
) -> None:
    files = {
        **_presentation("rId1"),
        "ppt/slides/slide1.xml": _slide(_shape("Kiến trúc", title=True), _picture("QR")),
        "ppt/slides/_rels/slide1.xml.rels": _rels(("rId9", "image", "../media/image1.tmp")),
    }
    media = LocalMediaStore(tmp_path / "media")

    result = _parse_deck(
        _deck(tmp_path, files, {"ppt/media/image1.tmp": JPEG}), media_store=media
    )

    figure = next(block for block in result.blocks if block.type is BlockType.FIGURE)
    assert figure.media_ref.endswith(".jpg")


def test_pptx_falls_back_to_a_slide_label_when_there_is_no_title(
    tmp_path: Path,
) -> None:
    files = {
        **_presentation("rId1"),
        "ppt/slides/slide1.xml": _slide(_picture("Ảnh nền")),
        "ppt/slides/_rels/slide1.xml.rels": _rels(("rId9", "image", "../media/image1.png")),
    }
    media = LocalMediaStore(tmp_path / "media")

    result = _parse_deck(
        _deck(tmp_path, files, {"ppt/media/image1.png": PNG}), media_store=media
    )

    assert result.blocks[0].type is BlockType.SLIDE
    assert result.blocks[0].text == "Slide 1"


def test_pptx_marks_a_hidden_slide_instead_of_dropping_it(tmp_path: Path) -> None:
    files = {
        **_presentation("rId1", "rId2"),
        "ppt/slides/slide1.xml": _slide(_shape("Doanh thu", title=True)),
        "ppt/slides/slide2.xml": _slide(_shape("Phụ lục", title=True), hidden=True),
    }

    result = _parse_deck(_deck(tmp_path, files))

    slides = [block for block in result.blocks if block.type is BlockType.SLIDE]
    assert [block.metadata.get("hidden") for block in slides] == [None, True]


def test_pptx_without_a_presentation_part_falls_back_to_part_names(
    tmp_path: Path,
) -> None:
    """A deck can still be read when the ordering part is missing or unreadable."""
    files = {
        "ppt/slides/slide1.xml": _slide(_shape("Mở đầu", title=True)),
        "ppt/slides/slide2.xml": _slide(_shape("Kết luận", title=True)),
    }

    result = _parse_deck(_deck(tmp_path, files))

    titles = [block.text for block in result.blocks if block.type is BlockType.SLIDE]
    assert titles == ["Mở đầu", "Kết luận"]


def test_pptx_carries_slide_text_and_speaker_notes(tmp_path: Path) -> None:
    files = {
        **_presentation("rId1"),
        "ppt/slides/slide1.xml": _slide(
            _shape("Quarterly review", title=True), _shape("Revenue increased")
        ),
        "ppt/slides/_rels/slide1.xml.rels": _rels(
            ("rId2", "notesSlide", "../notesSlides/notesSlide1.xml")
        ),
        "ppt/notesSlides/notesSlide1.xml": _notes("Explain the regional split."),
    }

    result = _parse_deck(_deck(tmp_path, files))

    assert result.document.source_type is SourceType.PPTX
    assert [block.type for block in result.blocks] == [
        BlockType.SLIDE,
        BlockType.PARAGRAPH,
        BlockType.SPEECH_TURN,
    ]
    assert result.blocks[1].text == "Revenue increased"
    assert "Explain the regional split." in result.blocks[2].text
    assert all(block.confidence == 1.0 for block in result.blocks)
