from __future__ import annotations

import html
import posixpath
import re
from datetime import date, datetime, time
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from rag_ingestion.models import (
    Block,
    BlockType,
    CellRangeLocator,
    Document,
    OfficePathLocator,
    ParseResult,
    SourceType,
)
from rag_ingestion.ingestion.parsers.datalab import ConvertedBlock, DocumentConverter
from rag_ingestion.logger import get_logger
from rag_ingestion.storage.media import MediaStore, media_key

_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_SHEET_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_PRESENTATION_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_SLIDE_SHAPE_TAGS = frozenset(
    f"{{{_PRESENTATION_NS}}}{tag}" for tag in ("sp", "pic", "graphicFrame")
)
_NS = {
    "w": _WORD_NS,
    "a": _DRAWING_NS,
    "pkg": _PKG_REL_NS,
    "xdr": _SHEET_DRAWING_NS,
}
_WORD_TAG = f"{{{_WORD_NS}}}"

_log = get_logger(__name__)


def _document(
    path: Path,
    *,
    owner_user_id: str,
    project_id: str,
    source_uri: str,
    source_type: SourceType,
) -> Document:
    return Document.from_content(
        content=path.read_bytes(),
        owner_user_id=owner_user_id,
        project_id=project_id,
        source_type=source_type,
        source_uri=source_uri,
        title=path.stem,
        parse_engine="native_ooxml",
        parse_version="1",
    )


def _word_text(element: ElementTree.Element) -> str:
    return "".join(node.text or "" for node in element.findall(".//w:t", _NS)).strip()


def parse_docx(
    path: Path, *, owner_user_id: str, project_id: str, source_uri: str
) -> ParseResult:
    document = _document(
        path,
        owner_user_id=owner_user_id,
        project_id=project_id,
        source_uri=source_uri,
        source_type=SourceType.DOCX,
    )
    with ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))

    body = root.find("w:body", _NS)
    if body is None:
        return ParseResult(document=document, blocks=[])

    blocks: list[Block] = []
    section_path: list[str] = []
    heading_ids: dict[int, str] = {}
    paragraph_index = 0
    table_index = 0
    for element in body:
        if element.tag == f"{_WORD_TAG}p":
            paragraph_index += 1
            text = _word_text(element)
            if not text:
                continue
            style_node = element.find("w:pPr/w:pStyle", _NS)
            style = style_node.get(f"{_WORD_TAG}val", "") if style_node is not None else ""
            heading_match = re.fullmatch(r"Heading\s*(\d+)", style, re.IGNORECASE)
            locator = OfficePathLocator(
                part="word/document.xml", path=f"/document/body/p[{paragraph_index}]"
            )
            if heading_match:
                level = max(1, int(heading_match.group(1)))
                section_path = section_path[: level - 1]
                section_path.append(text)
                parent_id = heading_ids.get(level - 1)
                block = Block.from_source(
                    doc_id=document.doc_id,
                    order=len(blocks),
                    type=BlockType.HEADING,
                    text=text,
                    confidence=1.0,
                    locators=[locator],
                    section_path=section_path.copy(),
                    parent_block_id=parent_id,
                    metadata={"style": style, "heading_level": level},
                )
                heading_ids = {key: value for key, value in heading_ids.items() if key < level}
                heading_ids[level] = block.block_id
            else:
                block = Block.from_source(
                    doc_id=document.doc_id,
                    order=len(blocks),
                    type=BlockType.PARAGRAPH,
                    text=text,
                    confidence=1.0,
                    locators=[locator],
                    section_path=section_path.copy(),
                    parent_block_id=heading_ids.get(max(heading_ids, default=0)),
                    metadata={"style": style or None},
                )
            blocks.append(block)
        elif element.tag == f"{_WORD_TAG}tbl":
            table_index += 1
            rows = [
                [_word_text(cell) for cell in row.findall("w:tc", _NS)]
                for row in element.findall("w:tr", _NS)
            ]
            table_html = "<table>" + "".join(
                "<tr>" + "".join(f"<td>{html.escape(cell)}</td>" for cell in row) + "</tr>"
                for row in rows
            ) + "</table>"
            table_locator = OfficePathLocator(
                part="word/document.xml", path=f"/document/body/tbl[{table_index}]"
            )
            table = Block.from_source(
                doc_id=document.doc_id,
                order=len(blocks),
                type=BlockType.TABLE,
                text=table_html,
                confidence=1.0,
                locators=[table_locator],
                section_path=section_path.copy(),
                parent_block_id=heading_ids.get(max(heading_ids, default=0)),
            )
            blocks.append(table)
            for row_number, cells in enumerate(rows, start=1):
                row_locator = OfficePathLocator(
                    part="word/document.xml",
                    path=f"/document/body/tbl[{table_index}]/tr[{row_number}]",
                )
                row = Block.from_source(
                    doc_id=document.doc_id,
                    order=len(blocks),
                    type=BlockType.TABLE_ROW,
                    text=" | ".join(cells),
                    confidence=1.0,
                    locators=[row_locator],
                    section_path=section_path.copy(),
                    parent_block_id=table.block_id,
                )
                blocks.append(row)
                for column_number, cell_text in enumerate(cells, start=1):
                    cell_locator = OfficePathLocator(
                        part="word/document.xml",
                        path=(
                            f"/document/body/tbl[{table_index}]/tr[{row_number}]"
                            f"/tc[{column_number}]"
                        ),
                    )
                    blocks.append(
                        Block.from_source(
                            doc_id=document.doc_id,
                            order=len(blocks),
                            type=BlockType.TABLE_CELL,
                            text=cell_text,
                            confidence=1.0,
                            locators=[cell_locator],
                            section_path=section_path.copy(),
                            parent_block_id=row.block_id,
                        )
                    )
    return ParseResult(document=document, blocks=blocks)


def _cell_text(value: object) -> str:
    """A cell as the text a reader would write, not as Python repr.

    Dates arrived as `2026-06-26 00:00:00` and whole numbers as `38.0`, neither
    of which appears anywhere in the spreadsheet the reader is looking at.
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:g}"
    return " ".join(str(value).split())


#: Leading bytes of every picture format Excel embeds. The stored extension is
#: not trustworthy -- workbooks ship JPEGs named `.tmp` -- and the content type
#: a downstream reader is handed has to match what the bytes actually are.
_IMAGE_MAGIC = (
    (b"\x89PNG\r\n\x1a\n", ".png"),
    (b"\xff\xd8\xff", ".jpg"),
    (b"GIF8", ".gif"),
    (b"BM", ".bmp"),
    (b"II*\x00", ".tif"),
    (b"MM\x00*", ".tif"),
)
_IMAGE_CONTENT_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".gif": "image/gif",
    ".bmp": "image/bmp",
    ".tif": "image/tiff",
    ".webp": "image/webp",
}


def _image_extension(payload: bytes) -> str:
    for magic, extension in _IMAGE_MAGIC:
        if payload.startswith(magic):
            return extension
    if payload[:4] == b"RIFF" and payload[8:12] == b"WEBP":
        return ".webp"
    return ".bin"


def _relationships(archive: ZipFile, part: str) -> dict[str, str]:
    """Relationship id to the zip entry it names, resolved to a full path."""
    base = posixpath.dirname(part)
    rels_part = posixpath.join(base, "_rels", f"{posixpath.basename(part)}.rels")
    if rels_part not in archive.namelist():
        return {}
    root = ElementTree.fromstring(archive.read(rels_part))
    resolved: dict[str, str] = {}
    for node in root.findall("pkg:Relationship", _NS):
        target = node.get("Target", "")
        if not target:
            continue
        resolved[node.get("Id", "")] = (
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.normpath(posixpath.join(base, target))
        )
    return resolved


def _anchor_cell(node: ElementTree.Element | None) -> str | None:
    if node is None:
        return None
    column = node.find("xdr:col", _NS)
    row = node.find("xdr:row", _NS)
    if column is None or row is None:
        return None
    return f"{get_column_letter(int(column.text or 0) + 1)}{int(row.text or 0) + 1}"


class _Picture:
    __slots__ = ("sheet", "name", "anchor", "payload", "extension")

    def __init__(self, *, sheet: str, name: str, anchor: str, payload: bytes) -> None:
        self.sheet = sheet
        self.name = name
        self.anchor = anchor
        self.payload = payload
        self.extension = _image_extension(payload)


def _sheet_pictures(path: Path) -> list[_Picture]:
    """Every embedded picture, with the sheet and cell range it covers.

    Read straight out of the package rather than through `Worksheet._images`,
    which openpyxl leaves empty unless Pillow is installed -- a missing
    optional dependency would otherwise drop every image without an error.
    """
    pictures: list[_Picture] = []
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        if "xl/workbook.xml" not in names:
            return pictures
        workbook_rels = _relationships(archive, "xl/workbook.xml")
        workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        for sheet_node in workbook.findall(".//{*}sheets/{*}sheet"):
            sheet_part = workbook_rels.get(sheet_node.get(f"{{{_DOC_REL_NS}}}id", ""))
            if not sheet_part or sheet_part not in names:
                continue
            title = sheet_node.get("name", "")
            for drawing_part in _relationships(archive, sheet_part).values():
                if not drawing_part.startswith("xl/drawings/") or drawing_part not in names:
                    continue
                drawing_rels = _relationships(archive, drawing_part)
                drawing = ElementTree.fromstring(archive.read(drawing_part))
                for position, anchor in enumerate(drawing, start=1):
                    picture = anchor.find("xdr:pic", _NS)
                    if picture is None:
                        continue
                    blip = picture.find("{*}blipFill/{*}blip")
                    if blip is None:
                        continue
                    media_part = drawing_rels.get(blip.get(f"{{{_DOC_REL_NS}}}embed", ""))
                    if not media_part or media_part not in names:
                        continue
                    start = _anchor_cell(anchor.find("xdr:from", _NS))
                    end = _anchor_cell(anchor.find("xdr:to", _NS))
                    if start is None:
                        continue
                    name_node = picture.find("xdr:nvPicPr/xdr:cNvPr", _NS)
                    label = name_node.get("name") if name_node is not None else None
                    pictures.append(
                        _Picture(
                            sheet=title,
                            name=label or f"Picture {position}",
                            anchor=start if end is None else f"{start}:{end}",
                            payload=archive.read(media_part),
                        )
                    )
    return pictures


class _Cell:
    __slots__ = ("row", "column", "text", "formula", "merged_range")

    def __init__(
        self,
        *,
        row: int,
        column: int,
        text: str,
        formula: str | None,
        merged_range: str | None,
    ) -> None:
        self.row = row
        self.column = column
        self.text = text
        self.formula = formula
        self.merged_range = merged_range

    @property
    def coordinate(self) -> str:
        return f"{get_column_letter(self.column)}{self.row}"


def _has_cached_results(workbook: Workbook, values: Workbook) -> bool:
    """Whether this workbook carries the results of its own formulas.

    The question has to be asked once for the whole file, never per cell. A
    calculated workbook stores nothing for a formula whose result was blank,
    so a per-cell fallback would print `=IF(I16=$K$14,J16,"")` in every cell
    that correctly evaluated to empty -- which is most of a real sheet.
    """
    for sheet in workbook.worksheets:
        value_sheet = values[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                if (
                    cell.data_type == "f"
                    and value_sheet.cell(cell.row, cell.column).value is not None
                ):
                    return True
    return False


def _sheet_rows(
    sheet: Worksheet, values: Worksheet, *, trust_cache: bool
) -> list[list[_Cell]]:
    """Rows carrying content, each cell keeping the column it was typed into.

    A cell shows its computed result. When the workbook was never calculated
    at all the formula text stands in, because a file Excel has not opened
    caches nothing and an empty sheet is worse than a visible formula. Rows
    holding only formulas that resolved to blank or zero are the fill-down
    tail below the data and are dropped.
    """
    merged = {
        (row, column)
        for item in sheet.merged_cells.ranges
        for row in range(item.min_row, item.max_row + 1)
        for column in range(item.min_col, item.max_col + 1)
    }
    anchors = {
        (item.min_row, item.min_col): str(item) for item in sheet.merged_cells.ranges
    }
    rows: list[list[_Cell]] = []
    for row_number in range(1, sheet.max_row + 1):
        cells: list[_Cell] = []
        has_content = False
        for column in range(1, sheet.max_column + 1):
            # Only the top-left cell of a merge holds the value, so rendering
            # the rest would repeat one caption across every column it spans.
            if (row_number, column) in merged and (row_number, column) not in anchors:
                continue
            source = sheet.cell(row_number, column)
            is_formula = source.data_type == "f"
            cached = values.cell(row_number, column).value
            text = _cell_text(
                source.value if is_formula and not trust_cache else cached
            )
            if not text:
                continue
            cells.append(
                _Cell(
                    row=row_number,
                    column=column,
                    text=text,
                    formula=str(source.value) if is_formula else None,
                    merged_range=anchors.get((row_number, column)),
                )
            )
            if not is_formula or text != "0":
                has_content = True
        if cells and has_content:
            rows.append(cells)
    return rows


def _grid_line(cells: dict[int, str], columns: list[int], label: str) -> str:
    body = " | ".join(cells.get(column, "") for column in columns)
    return f"| {label} | {body} |"


def parse_xlsx(
    path: Path,
    *,
    owner_user_id: str,
    project_id: str,
    source_uri: str,
    media_store: MediaStore | None = None,
) -> ParseResult:
    document = _document(
        path,
        owner_user_id=owner_user_id,
        project_id=project_id,
        source_uri=source_uri,
        source_type=SourceType.XLSX,
    )
    formula_workbook = load_workbook(path, data_only=False, read_only=False)
    value_workbook = load_workbook(path, data_only=True, read_only=False)
    # Without somewhere to put the bytes a figure block would carry a
    # `media_ref` that resolves to nothing, so the pictures are left alone.
    pictures_by_sheet: dict[str, list[_Picture]] = {}
    for picture in _sheet_pictures(path) if media_store is not None else []:
        pictures_by_sheet.setdefault(picture.sheet, []).append(picture)

    trust_cache = _has_cached_results(formula_workbook, value_workbook)
    blocks: list[Block] = []
    for sheet in formula_workbook.worksheets:
        rows = _sheet_rows(
            sheet, value_workbook[sheet.title], trust_cache=trust_cache
        )
        pictures = pictures_by_sheet.get(sheet.title, [])
        if not rows and not pictures:
            continue

        section_path = [sheet.title]
        sheet_range = (
            f"A1:{get_column_letter(max(1, sheet.max_column))}{max(1, sheet.max_row)}"
        )
        blocks.append(
            Block.from_source(
                doc_id=document.doc_id,
                order=len(blocks),
                type=BlockType.HEADING,
                text=sheet.title,
                confidence=1.0,
                locators=[CellRangeLocator(sheet=sheet.title, range=sheet_range)],
                section_path=section_path,
            )
        )

        if rows:
            columns = sorted({cell.column for cells in rows for cell in cells})
            letters = {column: get_column_letter(column) for column in columns}
            header = Block.from_source(
                doc_id=document.doc_id,
                order=len(blocks),
                type=BlockType.TABLE,
                text=(
                    _grid_line(letters, columns, "#")
                    + "\n|---|"
                    + "---|" * len(columns)
                ),
                confidence=1.0,
                locators=[CellRangeLocator(sheet=sheet.title, range=sheet_range)],
                section_path=section_path,
                metadata={"grid": sheet.title},
            )
            blocks.append(header)
            for cells in rows:
                row_number = cells[0].row
                row = Block.from_source(
                    doc_id=document.doc_id,
                    order=len(blocks),
                    type=BlockType.TABLE_ROW,
                    text=_grid_line(
                        {cell.column: cell.text.replace("|", "\\|") for cell in cells},
                        columns,
                        str(row_number),
                    ),
                    confidence=1.0,
                    locators=[
                        CellRangeLocator(
                            sheet=sheet.title,
                            range=(
                                f"{letters[columns[0]]}{row_number}:"
                                f"{letters[columns[-1]]}{row_number}"
                            ),
                        )
                    ],
                    parent_block_id=header.block_id,
                    section_path=section_path,
                    metadata={"grid": sheet.title, "row": row_number},
                )
                blocks.append(row)
                for cell in cells:
                    blocks.append(
                        Block.from_source(
                            doc_id=document.doc_id,
                            order=len(blocks),
                            type=BlockType.TABLE_CELL,
                            text=cell.text,
                            confidence=1.0,
                            locators=[
                                CellRangeLocator(
                                    sheet=sheet.title,
                                    range=cell.merged_range or cell.coordinate,
                                )
                            ],
                            parent_block_id=row.block_id,
                            section_path=section_path,
                            metadata={
                                "formula": cell.formula,
                                "merged_range": cell.merged_range,
                            },
                        )
                    )

        for picture in pictures:
            key = media_key(
                doc_id=document.doc_id,
                kind="figure",
                discriminator=f"{picture.sheet}!{picture.anchor}",
                extension=picture.extension,
            )
            stored = media_store.put(
                key,
                picture.payload,
                content_type=_IMAGE_CONTENT_TYPES.get(
                    picture.extension, "application/octet-stream"
                ),
            )
            blocks.append(
                Block.from_source(
                    doc_id=document.doc_id,
                    order=len(blocks),
                    type=BlockType.FIGURE,
                    text=f"{picture.name} ({picture.sheet}!{picture.anchor})",
                    confidence=1.0,
                    locators=[
                        CellRangeLocator(sheet=picture.sheet, range=picture.anchor)
                    ],
                    section_path=section_path,
                    media_ref=stored,
                )
            )

    formula_workbook.close()
    value_workbook.close()
    return ParseResult(document=document, blocks=blocks)


def _shape_paragraphs(shape: ElementTree.Element) -> list[str]:
    """The lines of one text frame, in the order they are written.

    Read per paragraph rather than by sweeping every `a:t` on the slide, so a
    bullet stays a line of its own and the chunker can size and split the body
    the way it does for any other prose.
    """
    lines: list[str] = []
    for paragraph in shape.iter(f"{{{_DRAWING_NS}}}p"):
        # A field holds the slide number or the date PowerPoint fills in, not
        # anything the author wrote.
        generated = {
            id(node)
            for field in paragraph.iter(f"{{{_DRAWING_NS}}}fld")
            for node in field.iter(f"{{{_DRAWING_NS}}}t")
        }
        text = " ".join(
            piece
            for node in paragraph.iter(f"{{{_DRAWING_NS}}}t")
            if id(node) not in generated and (piece := (node.text or "").strip())
        )
        if text:
            lines.append(text)
    return lines


def _is_title(shape: ElementTree.Element) -> bool:
    return any(
        placeholder.get("type") in {"title", "ctrTitle"}
        for placeholder in shape.iter(f"{{{_PRESENTATION_NS}}}ph")
    )


def _table_rows(frame: ElementTree.Element) -> list[list[str]]:
    table = next(iter(frame.iter(f"{{{_DRAWING_NS}}}tbl")), None)
    if table is None:
        return []
    return [
        [" ".join(_shape_paragraphs(cell)) for cell in row.iter(f"{{{_DRAWING_NS}}}tc")]
        for row in table.iter(f"{{{_DRAWING_NS}}}tr")
    ]


def _slide_parts(archive: ZipFile, names: set[str]) -> list[str]:
    """Slide parts in presentation order.

    Reordering a deck rewrites `sldIdLst` and never renames a part, so sorting
    by file name shows the reader a different deck than the author built. The
    part names are the fallback for a package missing its presentation part.
    """
    if "ppt/presentation.xml" in names:
        relationships = _relationships(archive, "ppt/presentation.xml")
        root = ElementTree.fromstring(archive.read("ppt/presentation.xml"))
        ordered = [
            part
            for node in root.iter(f"{{{_PRESENTATION_NS}}}sldId")
            if (part := relationships.get(node.get(f"{{{_DOC_REL_NS}}}id", "")))
            in names
        ]
        if ordered:
            return ordered
    return sorted(
        (name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
        key=lambda name: int(re.search(r"\d+", Path(name).stem).group()),
    )


def parse_pptx(
    path: Path,
    *,
    owner_user_id: str,
    project_id: str,
    source_uri: str,
    media_store: MediaStore | None = None,
    converter: DocumentConverter | None = None,
) -> ParseResult:
    """Read a deck, by rendering it when that is possible.

    The shape tree is all text and no picture, and a deck is mostly picture: a
    diagram drawn as shapes says nothing about which box feeds which, and one
    pasted in as a flat image says nothing at all. A converter renders each
    slide and reads it, which is the only way either of those becomes
    searchable. It is also a paid network call, so the shape tree stays as the
    floor for when the service is unavailable.
    """
    if converter is not None:
        try:
            converted = converter.convert(path)
        except Exception:
            _log.exception(
                "Slide conversion failed for %s; falling back to the shape tree",
                path.name,
            )
            converted = []
        if converted:
            return _converted_result(
                converted,
                path=path,
                owner_user_id=owner_user_id,
                project_id=project_id,
                source_uri=source_uri,
                media_store=media_store,
            )
        _log.warning(
            "Slide conversion returned nothing for %s; falling back to the shape tree",
            path.name,
        )
    document = _document(
        path,
        owner_user_id=owner_user_id,
        project_id=project_id,
        source_uri=source_uri,
        source_type=SourceType.PPTX,
    )
    blocks: list[Block] = []
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        for slide_number, part in enumerate(_slide_parts(archive, names), start=1):
            blocks.extend(
                _slide_blocks(
                    archive,
                    names,
                    doc_id=document.doc_id,
                    part=part,
                    slide_number=slide_number,
                    order=len(blocks),
                    media_store=media_store,
                )
            )
    return ParseResult(document=document, blocks=blocks)


#: How many slides a picture has to appear on before it is read as decoration.
#: Matches the threshold `reconciliation` uses to drop a running page header.
_REPEATED_FIGURE_SLIDES = 3


def _repeated_figures(converted: list[ConvertedBlock]) -> set[str]:
    """Captions belonging to a picture the master slide puts on every page.

    Having a caption is not enough to make a picture content: a logo is
    described as readily as a diagram, and one repeated across a deck was a
    quarter of the chunks it produced.
    """
    slides_by_caption: dict[str, set[int]] = {}
    for block in converted:
        if block.block_type is BlockType.FIGURE and block.text:
            slides_by_caption.setdefault(block.text, set()).add(block.page)
    return {
        caption
        for caption, slides in slides_by_caption.items()
        if len(slides) >= _REPEATED_FIGURE_SLIDES
    }


def _converted_result(
    converted: list[ConvertedBlock],
    *,
    path: Path,
    owner_user_id: str,
    project_id: str,
    source_uri: str,
    media_store: MediaStore | None,
) -> ParseResult:
    document = _document(
        path,
        owner_user_id=owner_user_id,
        project_id=project_id,
        source_uri=source_uri,
        source_type=SourceType.PPTX,
    )
    for block in converted:
        if block.page < 1:
            raise ValueError(f"Conversion returned an impossible slide number: {block.page}")

    repeated = _repeated_figures(converted)
    by_slide: dict[int, list[ConvertedBlock]] = {}
    for block in converted:
        if block.furniture:
            continue
        if block.block_type is BlockType.FIGURE and block.text in repeated:
            continue
        by_slide.setdefault(block.page, []).append(block)

    blocks: list[Block] = []
    for slide_number in sorted(by_slide):
        contents = by_slide[slide_number]
        part = f"ppt/slides/slide{slide_number}.xml"
        title = next(
            (
                item.text
                for item in contents
                if item.block_type is BlockType.HEADING and item.text
            ),
            None,
        )
        section_path = [title or f"Slide {slide_number}"]

        def add(
            type: BlockType,
            text: str,
            *,
            path_suffix: str,
            media_ref: str | None = None,
            parent_block_id: str | None = None,
        ) -> Block:
            block = Block.from_source(
                doc_id=document.doc_id,
                order=len(blocks),
                type=type,
                text=text,
                confidence=1.0,
                locators=[OfficePathLocator(part=part, path=path_suffix)],
                section_path=section_path.copy(),
                parent_block_id=parent_block_id,
                media_ref=media_ref,
                metadata={"slide": slide_number},
            )
            blocks.append(block)
            return block

        slide = add(BlockType.SLIDE, section_path[0], path_suffix="/sld/cSld")
        used_title = title is not None
        for position, item in enumerate(contents, start=1):
            if used_title and item.block_type is BlockType.HEADING and item.text == title:
                used_title = False
                continue
            if item.block_type is BlockType.FIGURE:
                # A logo repeated on every slide comes back with no caption,
                # and eighteen copies of it are eighteen chunks of nothing.
                if not item.text:
                    continue
                reference = None
                if media_store is not None and item.media is not None:
                    extension = _image_extension(item.media)
                    reference = media_store.put(
                        media_key(
                            doc_id=document.doc_id,
                            kind="figure",
                            discriminator=f"{part}:{position}",
                            extension=extension,
                        ),
                        item.media,
                        content_type=_IMAGE_CONTENT_TYPES.get(
                            extension, "application/octet-stream"
                        ),
                    )
                add(
                    BlockType.FIGURE,
                    item.text,
                    path_suffix=f"/sld/cSld/spTree/pic[{position}]",
                    media_ref=reference,
                    parent_block_id=slide.block_id,
                )
                continue
            if not item.text:
                continue
            add(
                item.block_type,
                item.text,
                path_suffix=f"/sld/cSld/spTree/sp[{position}]",
                parent_block_id=slide.block_id,
            )
    return ParseResult(document=document, blocks=blocks)


def _slide_blocks(
    archive: ZipFile,
    names: set[str],
    *,
    doc_id: str,
    part: str,
    slide_number: int,
    order: int,
    media_store: MediaStore | None,
) -> list[Block]:
    root = ElementTree.fromstring(archive.read(part))
    tree = next(iter(root.iter(f"{{{_PRESENTATION_NS}}}spTree")), None)
    relationships = _relationships(archive, part)
    shapes = _flatten_shapes(tree)

    title, consumed = _slide_title(shapes)
    section_path = [title or f"Slide {slide_number}"]

    blocks: list[Block] = []

    def add(
        type: BlockType,
        text: str,
        *,
        path: str,
        parent_block_id: str | None = None,
        media_ref: str | None = None,
        **metadata: object,
    ) -> Block:
        block = Block.from_source(
            doc_id=doc_id,
            order=order + len(blocks),
            type=type,
            text=text,
            confidence=1.0,
            locators=[OfficePathLocator(part=part, path=path)],
            section_path=section_path.copy(),
            parent_block_id=parent_block_id,
            media_ref=media_ref,
            metadata={"slide": slide_number, **metadata},
        )
        blocks.append(block)
        return block

    slide = add(
        BlockType.SLIDE,
        section_path[0],
        path="/sld/cSld",
        **({"hidden": True} if root.get("show") == "0" else {}),
    )

    shape_index = table_index = picture_index = 0
    for position, shape in enumerate(shapes):
        if shape.tag == f"{{{_PRESENTATION_NS}}}sp":
            shape_index += 1
            for line_number, line in enumerate(_shape_paragraphs(shape), start=1):
                if (position, line_number) == consumed:
                    continue
                add(
                    BlockType.PARAGRAPH,
                    line,
                    path=f"/sld/cSld/spTree/sp[{shape_index}]/txBody/p[{line_number}]",
                    parent_block_id=slide.block_id,
                )
        elif shape.tag == f"{{{_PRESENTATION_NS}}}graphicFrame":
            rows = [row for row in _table_rows(shape) if any(cell for cell in row)]
            if not rows:
                continue
            table_index += 1
            header, *body = rows
            path = f"/sld/cSld/spTree/graphicFrame[{table_index}]"
            grid = f"{part}:{table_index}"
            container = add(
                BlockType.TABLE,
                _row_line(header) + "\n|" + "---|" * len(header),
                path=path,
                parent_block_id=slide.block_id,
                grid=grid,
            )
            for row_number, cells in enumerate(body, start=2):
                add(
                    BlockType.TABLE_ROW,
                    _row_line(cells),
                    path=f"{path}/tr[{row_number}]",
                    parent_block_id=container.block_id,
                    grid=grid,
                    row=row_number,
                )
        elif shape.tag == f"{{{_PRESENTATION_NS}}}pic" and media_store is not None:
            picture_index += 1
            stored = _store_slide_picture(
                archive,
                names,
                relationships,
                shape,
                doc_id=doc_id,
                part=part,
                picture_index=picture_index,
                media_store=media_store,
            )
            if stored is None:
                continue
            name, reference = stored
            add(
                BlockType.FIGURE,
                f"{name} (slide {slide_number})",
                path=f"/sld/cSld/spTree/pic[{picture_index}]",
                parent_block_id=slide.block_id,
                media_ref=reference,
            )

    notes = _notes_text(archive, names, relationships)
    if notes:
        # Its own block, and its own chunk: a reader could not tell the spoken
        # script from what was on the screen when both shared one block.
        add(
            BlockType.SPEECH_TURN,
            f"Speaker notes: {notes}",
            path="/notes/cSld",
            parent_block_id=slide.block_id,
            speaker_notes=True,
        )
    return blocks


def _row_line(cells: list[str]) -> str:
    return "| " + " | ".join(cell.replace("|", "\\|") for cell in cells) + " |"


def _flatten_shapes(tree: ElementTree.Element | None) -> list[ElementTree.Element]:
    """Every shape on a slide in drawing order, groups opened up.

    Grouping is how a deck draws a diagram -- select the boxes, Ctrl+G -- so
    reading only the top level of the tree drops exactly the slides that carry
    the most meaning.
    """
    if tree is None:
        return []
    shapes: list[ElementTree.Element] = []
    for shape in tree:
        if shape.tag == f"{{{_PRESENTATION_NS}}}grpSp":
            shapes.extend(_flatten_shapes(shape))
        elif shape.tag in _SLIDE_SHAPE_TAGS:
            shapes.append(shape)
    return shapes


def _slide_title(shapes: list[ElementTree.Element]) -> tuple[str | None, tuple[int, int] | None]:
    """The slide's title, and the body line it was taken from if it was borrowed.

    A title placeholder is the answer when there is one. Otherwise the first
    line of text on the slide is the closest thing to a title, and it is
    returned along with its position so the body does not repeat it.
    """
    for position, shape in enumerate(shapes):
        if shape.tag == f"{{{_PRESENTATION_NS}}}sp" and _is_title(shape):
            lines = _shape_paragraphs(shape)
            if lines:
                return lines[0], (position, 1)
    for position, shape in enumerate(shapes):
        if shape.tag == f"{{{_PRESENTATION_NS}}}sp":
            lines = _shape_paragraphs(shape)
            if lines:
                return lines[0], (position, 1)
    return None, None


def _store_slide_picture(
    archive: ZipFile,
    names: set[str],
    relationships: dict[str, str],
    shape: ElementTree.Element,
    *,
    doc_id: str,
    part: str,
    picture_index: int,
    media_store: MediaStore,
) -> tuple[str, str] | None:
    blip = next(iter(shape.iter(f"{{{_DRAWING_NS}}}blip")), None)
    if blip is None:
        return None
    media_part = relationships.get(blip.get(f"{{{_DOC_REL_NS}}}embed", ""))
    if media_part is None or media_part not in names:
        return None
    properties = next(iter(shape.iter(f"{{{_PRESENTATION_NS}}}cNvPr")), None)
    name = (properties.get("name") if properties is not None else None) or (
        f"Picture {picture_index}"
    )
    payload = archive.read(media_part)
    extension = _image_extension(payload)
    reference = media_store.put(
        media_key(
            doc_id=doc_id,
            kind="figure",
            discriminator=f"{part}:{picture_index}",
            extension=extension,
        ),
        payload,
        content_type=_IMAGE_CONTENT_TYPES.get(extension, "application/octet-stream"),
    )
    return name, reference


def _notes_text(
    archive: ZipFile, names: set[str], relationships: dict[str, str]
) -> str:
    """The speaker notes a slide actually owns.

    Resolved through the slide's own relationships. The part number is
    creation order, not slide order, so guessing `notesSlide{n}.xml` handed
    one slide's script to another and read as though the author wrote it there.
    """
    part = next(
        (
            target
            for target in relationships.values()
            if target.startswith("ppt/notesSlides/") and target in names
        ),
        None,
    )
    if part is None:
        return ""
    root = ElementTree.fromstring(archive.read(part))
    tree = next(iter(root.iter(f"{{{_PRESENTATION_NS}}}spTree")), None)
    if tree is None:
        return ""
    lines = [
        line
        for shape in tree.iter(f"{{{_PRESENTATION_NS}}}sp")
        for line in _shape_paragraphs(shape)
    ]
    return "\n".join(lines)
